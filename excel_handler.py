from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from werkzeug.security import generate_password_hash, check_password_hash
import os
from datetime import datetime

class ExcelHandler:
    def __init__(self, filename='game_data.xlsx'):
        self.filename = filename
    
    def initialize_excel(self):
        """Create Excel file with sheets if it doesn't exist"""
        if not os.path.exists(self.filename):
            wb = Workbook()
            
            # Users sheet
            ws_users = wb.active
            ws_users.title = 'Users'
            ws_users['A1'] = 'Username'
            ws_users['B1'] = 'Password_Hash'
            ws_users['C1'] = 'Joined_Date'
            
            # Predictions sheet
            ws_predictions = wb.create_sheet('Predictions')
            ws_predictions['A1'] = 'Username'
            ws_predictions['B1'] = 'Week'
            ws_predictions['C1'] = 'Price_Guess'
            ws_predictions['D1'] = 'Submitted_Date'
            
            # Results sheet
            ws_results = wb.create_sheet('Results')
            ws_results['A1'] = 'Week'
            ws_results['B1'] = 'Actual_Price'
            ws_results['C1'] = 'Updated_Date'
            
            wb.save(self.filename)
    
    def add_user(self, username, password):
        """Add a new user"""
        wb = load_workbook(self.filename)
        ws = wb['Users']
        
        row = ws.max_row + 1
        ws[f'A{row}'] = username
        ws[f'B{row}'] = generate_password_hash(password)
        ws[f'C{row}'] = datetime.now().isoformat()
        
        wb.save(self.filename)
    
    def user_exists(self, username):
        """Check if user exists"""
        wb = load_workbook(self.filename)
        ws = wb['Users']
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == username:
                return True
        return False
    
    def verify_user(self, username, password):
        """Verify user credentials"""
        wb = load_workbook(self.filename)
        ws = wb['Users']
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == username:
                return check_password_hash(row[1], password)
        return False
    
    def add_prediction(self, username, week, price_guess):
        """Add a prediction for the week"""
        wb = load_workbook(self.filename)
        ws = wb['Predictions']
        
        # Check if user already submitted for this week
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == username and row[1] == week:
                # Update existing prediction
                for cell_row in ws.iter_rows(min_row=2):
                    if cell_row[0].value == username and cell_row[1].value == week:
                        cell_row[2].value = price_guess
                        cell_row[3].value = datetime.now().isoformat()
                wb.save(self.filename)
                return
        
        # Add new prediction
        row = ws.max_row + 1
        ws[f'A{row}'] = username
        ws[f'B{row}'] = week
        ws[f'C{row}'] = price_guess
        ws[f'D{row}'] = datetime.now().isoformat()
        
        wb.save(self.filename)
    
    def has_submitted_this_week(self, username):
        """Check if user has submitted prediction for current week"""
        from utils import get_week_number
        wb = load_workbook(self.filename)
        ws = wb['Predictions']
        current_week = get_week_number()
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == username and row[1] == current_week:
                return True
        return False
    
    def get_week_predictions(self, week):
        """Get all predictions for a week"""
        wb = load_workbook(self.filename)
        ws = wb['Predictions']
        
        predictions = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[1] == week:
                predictions.append({
                    'username': row[0],
                    'week': row[1],
                    'guess': row[2],
                    'submitted': row[3]
                })
        return predictions
    
    def set_actual_price(self, week, actual_price):
        """Set the actual end-of-week price"""
        wb = load_workbook(self.filename)
        ws = wb['Results']
        
        # Check if week already has a result
        for cell_row in ws.iter_rows(min_row=2):
            if cell_row[0].value == week:
                cell_row[1].value = actual_price
                cell_row[2].value = datetime.now().isoformat()
                wb.save(self.filename)
                return
        
        # Add new result
        row = ws.max_row + 1
        ws[f'A{row}'] = week
        ws[f'B{row}'] = actual_price
        ws[f'C{row}'] = datetime.now().isoformat()
        
        wb.save(self.filename)
    
    def get_actual_price(self, week):
        """Get actual price for a week"""
        wb = load_workbook(self.filename)
        ws = wb['Results']
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] == week:
                return row[1]
        return None
    
    def get_user_stats(self, username):
        """Get user's cumulative stats"""
        wb = load_workbook(self.filename)
        ws_pred = wb['Predictions']
        ws_results = wb['Results']
        
        total_points = 0
        weeks_participated = 0
        
        for pred_row in ws_pred.iter_rows(min_row=2, values_only=True):
            if pred_row[0] == username:
                week = pred_row[1]
                guess = pred_row[2]
                
                actual_price = None
                for result_row in ws_results.iter_rows(min_row=2, values_only=True):
                    if result_row[0] == week:
                        actual_price = result_row[1]
                        break
                
                if actual_price is not None:
                    from scoring import calculate_points
                    points = calculate_points(guess, actual_price)
                    total_points += points
                    weeks_participated += 1
        
        return {
            'total_points': total_points,
            'weeks_participated': weeks_participated
        }
    
    def get_leaderboard(self):
        """Get current leaderboard"""
        wb = load_workbook(self.filename)
        ws_users = wb['Users']
        
        scores = {}
        
        for user_row in ws_users.iter_rows(min_row=2, values_only=True):
            username = user_row[0]
            stats = self.get_user_stats(username)
            scores[username] = stats['total_points']
        
        # Sort by points
        leaderboard = sorted(scores.items(), key=lambda x: x[1], reverse=True)
        return [{'rank': i+1, 'username': name, 'points': points} 
                for i, (name, points) in enumerate(leaderboard)]
